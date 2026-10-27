# admin.py
#!/usr/bin/env python3
"""
Flask Blueprint для управления заказами (продукты, матрасы), клиентами,
галереей и конструктором матрасов в админ-панели.

Функциональность матрасов (слои, чехлы, размеры, цены, характеристики,
фото, экспорт Excel) перенесена из Telegram-бота toppers_bot.py.

Раздел «Партнёрские цены» (аналог partners_prices_bot.py):
  * список файлов цен партнёров в каталоге partners/,
  * генерация/скачивание шаблона 000000.xlsx из БД,
  * создание копии из 000000.xlsx,
  * изменение всех цен на процент,
  * изменение одной цены (set/mul/add/sub),
  * переименование и удаление файлов,
  * загрузка своего xlsx с валидацией структуры.
"""

import io
import json
import logging
import os
import re
import shutil
import tempfile
from decimal import Decimal
from datetime import datetime, date, time
from pathlib import Path

import psycopg2
from flask import Blueprint, request, session, send_file
from flask import jsonify as _flask_jsonify
from werkzeug.security import generate_password_hash

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:  # pragma: no cover
    Image = None
    PIL_AVAILABLE = False

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover
    Workbook = None
    load_workbook = None
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger(__name__)


# ------------------------- JSON-сериализация -------------------------
def _json_safe(obj):
    if isinstance(obj, Decimal):
        if obj == obj.to_integral_value():
            return int(obj)
        return float(obj)
    if isinstance(obj, datetime):
        return obj.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(obj, date):
        return obj.strftime('%Y-%m-%d')
    if isinstance(obj, time):
        return obj.strftime('%H:%M:%S')
    if isinstance(obj, dict):
        return {k: _json_safe(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_json_safe(v) for v in obj]
    return obj


def jsonify(data):
    return _flask_jsonify(_json_safe(data))


# ------------------------- Конфигурация -------------------------
BASE_DIR = Path(__file__).resolve().parent

PAGE_SIZE = 8
CLIENTS_PAGE_SIZE = 8
CLIENT_ORDERS_PAGE_SIZE = 5
GALLERY_PAGE_SIZE = 8
MATTRESS_PAGE_SIZE = 8

STATUSES = [
    "Ожидает подтверждения",
    "ожидает оплаты",
    "в сборке",
    "в пути",
    "доставлен",
]

MEDIA_BASE = BASE_DIR.parent / "media" if (BASE_DIR.parent / "media").exists() else BASE_DIR / "media"
GALLERY_BASE = MEDIA_BASE / "gallery"
COMMERCIAL_BASE = MEDIA_BASE / "commercial"
LAYERS_MEDIA = MEDIA_BASE / "layers"
COVERS_MEDIA = MEDIA_BASE / "covers"

for _d in (GALLERY_BASE, COMMERCIAL_BASE, LAYERS_MEDIA, COVERS_MEDIA):
    _d.mkdir(parents=True, exist_ok=True)

# Партнёрские цены — та же логика поиска каталога, что и у media
PARTNERS_BASE = BASE_DIR.parent / "partners" if (BASE_DIR.parent / "partners").exists() else BASE_DIR / "partners"
PARTNERS_BASE.mkdir(parents=True, exist_ok=True)
ORIGINAL_PARTNER_NAME = "000000"

GALLERY_SECTIONS = {
    'gallery':    {'base': GALLERY_BASE,    'title': 'Индивидуальная галерея'},
    'commercial': {'base': COMMERCIAL_BASE, 'title': 'Коммерческая галерея'},
}


def _read_db_config():
    candidates = [
        BASE_DIR / 'config_db.txt',
        BASE_DIR / 'telegram_admin_panel' / 'config_db.txt',
        BASE_DIR.parent / 'config_db.txt',
        BASE_DIR.parent / 'telegram_admin_panel' / 'config_db.txt',
    ]
    for config_path in candidates:
        if not config_path.exists():
            continue
        config = {}
        with open(config_path, 'r', encoding='utf-8') as f:
            for line in f:
                if '=' in line:
                    k, v = line.strip().split('=', 1)
                    config[k.strip()] = v.strip()
        required = ('database', 'user', 'password')
        if all(key in config for key in required):
            logger.info(f"admin.py: использую конфиг БД {config_path}")
            return config
        logger.warning(f"admin.py: в {config_path} не хватает обязательных ключей {required}")
    raise FileNotFoundError(
        "config_db.txt не найден ни по одному из путей: "
        + ", ".join(str(p) for p in candidates)
    )


_DB_CONFIG = _read_db_config()


def _get_db_connection():
    return psycopg2.connect(
        host=_DB_CONFIG.get('host', 'localhost'),
        port=_DB_CONFIG.get('port', '5432'),
        database=_DB_CONFIG['database'],
        user=_DB_CONFIG['user'],
        password=_DB_CONFIG['password'],
    )


# ------------------------- Blueprint -------------------------
admin_bp = Blueprint('admin_orders', __name__)


def _require_admin():
    return bool(session.get('admin_authenticated'))


def _unauthorized():
    return jsonify({'error': 'Не авторизован'}), 401


def _format_dt(dt):
    if not dt:
        return None
    return dt.strftime('%Y-%m-%d %H:%M:%S')


def _format_dt_ru(dt):
    if not dt:
        return None
    return dt.strftime('%d.%m.%Y %H:%M')


def _json_body():
    data = request.get_json(force=True, silent=True)
    if data is None:
        data = {}
    return data


def _table_exists(cur, table_name):
    cur.execute(
        "SELECT EXISTS(SELECT 1 FROM information_schema.tables "
        "WHERE table_schema='public' AND table_name=%s)",
        (table_name,),
    )
    return bool(cur.fetchone()[0])


def _table_columns(cur, table_name):
    cur.execute(
        "SELECT column_name FROM information_schema.columns "
        "WHERE table_schema='public' AND table_name=%s",
        (table_name,),
    )
    return {r[0] for r in cur.fetchall()}


# ============================================================
#                    СТАТУСЫ
# ============================================================
@admin_bp.route('/api/admin/orders/statuses')
def get_statuses():
    if not _require_admin():
        return _unauthorized()
    return jsonify({'statuses': STATUSES})


# ============================================================
#                    ЗАКАЗЫ ПРОДУКТОВ
# ============================================================
@admin_bp.route('/api/admin/orders/products')
def list_product_orders():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM orders")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * PAGE_SIZE

        cur.execute(
            "SELECT id, user_name, phone, order_date, status FROM orders "
            "ORDER BY order_date DESC OFFSET %s LIMIT %s",
            (offset, PAGE_SIZE),
        )
        orders = [{
            'id': r[0],
            'user_name': r[1],
            'phone': r[2],
            'order_date': _format_dt(r[3]),
            'status': r[4],
        } for r in cur.fetchall()]

        return jsonify({
            'orders': orders,
            'total': total,
            'page': page,
            'page_size': PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_product_orders")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/orders/products/<int:order_id>')
def product_order_details(order_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT id, user_name, phone, email, address, comment,
                   contact_time, order_date, status
            FROM orders WHERE id = %s
        """, (order_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Заказ не найден'}), 404

        order = {
            'id': r[0],
            'user_name': r[1],
            'phone': r[2],
            'email': r[3],
            'address': r[4],
            'comment': r[5],
            'contact_time': r[6],
            'order_date': _format_dt(r[7]),
            'status': r[8],
        }

        cur.execute("""
            SELECT COUNT(*) FROM information_schema.columns
            WHERE table_name='order_items' AND column_name='extra_data'
        """)
        has_extra = cur.fetchone()[0] > 0

        if has_extra:
            cur.execute("""
                SELECT product_code, product_name, material_code, material_name,
                       cost, quantity, extra_data
                FROM order_items WHERE order_id = %s ORDER BY id
            """, (order_id,))
        else:
            cur.execute("""
                SELECT product_code, product_name, material_code, material_name,
                       cost, quantity, NULL::text AS extra_data
                FROM order_items WHERE order_id = %s ORDER BY id
            """, (order_id,))

        items = []
        total_sum = 0
        for it in cur.fetchall():
            qty = it[5] if it[5] else 1
            item_total = (it[4] or 0) * qty
            total_sum += item_total
            extra = None
            is_topper = False
            if it[6]:
                if isinstance(it[6], dict):
                    extra = it[6]
                    is_topper = True
                else:
                    try:
                        extra = json.loads(it[6])
                        is_topper = True
                    except (ValueError, TypeError):
                        extra = None
            items.append({
                'product_code': it[0],
                'product_name': it[1],
                'material_code': it[2],
                'material_name': it[3],
                'cost': it[4],
                'quantity': qty,
                'total': item_total,
                'extra_data': extra,
                'is_topper': is_topper,
            })

        return jsonify({
            'order': order,
            'items': items,
            'total_sum': total_sum,
        })
    except Exception as e:
        logger.exception("admin.py: product_order_details")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/orders/products/<int:order_id>/status', methods=['PUT'])
def update_product_status(order_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    new_status = (data.get('status') or '').strip()
    if new_status not in STATUSES:
        return jsonify({'error': 'Недопустимый статус'}), 400
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("UPDATE orders SET status = %s WHERE id = %s", (new_status, order_id))
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Заказ не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_product_status")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/orders/products/<int:order_id>', methods=['DELETE'])
def delete_product_order(order_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM orders WHERE id = %s", (order_id,))
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Заказ не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_product_order")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#                    ЗАКАЗЫ МАТРАСОВ
# ============================================================
@admin_bp.route('/api/admin/orders/mattresses')
def list_mattress_orders():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    empty = {
        'orders': [], 'total': 0, 'page': 0,
        'page_size': PAGE_SIZE, 'total_pages': 1,
    }
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if not _table_exists(cur, 'mattress_orders'):
            return jsonify(empty)

        cols = _table_columns(cur, 'mattress_orders')
        size_join = ""
        size_select = "'—' AS size_name"
        if 'size_id' in cols and _table_exists(cur, 'mattress_sizes'):
            size_join = "LEFT JOIN mattress_sizes s ON o.size_id = s.id"
            size_select = "COALESCE(s.name, '—') AS size_name"

        date_col = 'order_date' if 'order_date' in cols else 'id'
        order_dir = 'DESC' if 'order_date' in cols else 'DESC'

        cur.execute("SELECT COUNT(*) FROM mattress_orders")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * PAGE_SIZE

        cur.execute(f"""
            SELECT o.id, o.user_name, o.phone, o.{date_col}, o.status,
                   {size_select}
            FROM mattress_orders o
            {size_join}
            ORDER BY o.{date_col} {order_dir}
            OFFSET %s LIMIT %s
        """, (offset, PAGE_SIZE))
        rows = cur.fetchall()
        orders = [{
            'id': r[0],
            'user_name': r[1],
            'phone': r[2],
            'order_date': _format_dt(r[3]),
            'status': r[4],
            'size_name': r[5],
        } for r in rows]

        return jsonify({
            'orders': orders,
            'total': total,
            'page': page,
            'page_size': PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_mattress_orders")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/orders/mattresses/<int:order_id>')
def mattress_order_details(order_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if not _table_exists(cur, 'mattress_orders'):
            return jsonify({'error': 'Заказ не найден'}), 404

        cols = _table_columns(cur, 'mattress_orders')

        base_cols = ['id', 'user_name', 'phone', 'email', 'address',
                     'comment', 'contact_time', 'order_date', 'status']
        select_parts = []
        for c in base_cols:
            if c in cols:
                select_parts.append(c)
            else:
                select_parts.append(f"NULL AS {c}")
        for c in ['size_id', 'cover_id', 'initial_height', 'total_height',
                  'total_price', 'layers_height']:
            if c in cols:
                select_parts.append(c)
            else:
                select_parts.append(f"NULL AS {c}")

        cur.execute(
            f"SELECT {', '.join(select_parts)} FROM mattress_orders WHERE id = %s",
            (order_id,),
        )
        row = cur.fetchone()
        if not row:
            return jsonify({'error': 'Заказ не найден'}), 404

        col_names = base_cols + ['size_id', 'cover_id', 'initial_height',
                                 'total_height', 'total_price', 'layers_height']
        o = dict(zip(col_names, row))

        size_name = '—'
        if o.get('size_id') and _table_exists(cur, 'mattress_sizes'):
            cur.execute("SELECT name FROM mattress_sizes WHERE id = %s", (o['size_id'],))
            r = cur.fetchone()
            if r:
                size_name = r[0]

        cover_name = '—'
        cover_price = 0
        if o.get('cover_id') and _table_exists(cur, 'mattress_cover'):
            cur.execute("SELECT name FROM mattress_cover WHERE id = %s", (o['cover_id'],))
            r = cur.fetchone()
            if r:
                cover_name = r[0]
            if o.get('size_id') and _table_exists(cur, 'mattress_cover_prices'):
                cur.execute(
                    "SELECT price FROM mattress_cover_prices "
                    "WHERE cover_id = %s AND size_id = %s",
                    (o['cover_id'], o['size_id']),
                )
                r = cur.fetchone()
                if r and r[0] is not None:
                    cover_price = r[0]

        order = {
            'id': o['id'],
            'user_name': o['user_name'],
            'phone': o['phone'],
            'email': o['email'],
            'address': o['address'],
            'comment': o['comment'],
            'contact_time': o['contact_time'],
            'order_date': _format_dt(o['order_date']),
            'status': o['status'],
            'size_name': size_name,
            'initial_height': o.get('initial_height') or 0,
            'cover_name': cover_name,
            'cover_price': cover_price,
        }

        items = []
        items_total = 0
        if _table_exists(cur, 'mattress_order_layers'):
            item_cols = _table_columns(cur, 'mattress_order_layers')
            item_select = []
            for c in ['layer_id', 'quantity', 'price_per_unit', 'total_price']:
                item_select.append(c if c in item_cols else f"NULL AS {c}")

            cur.execute(
                f"SELECT {', '.join(item_select)} FROM mattress_order_layers "
                f"WHERE order_id = %s ORDER BY id",
                (order_id,),
            )
            for r in cur.fetchall():
                layer_id, qty, price_per_unit, total_price = r
                layer_name = '—'
                layer_code = '—'
                if layer_id and _table_exists(cur, 'mattress_layers'):
                    cur.execute(
                        "SELECT code, name FROM mattress_layers WHERE id = %s",
                        (layer_id,),
                    )
                    lr = cur.fetchone()
                    if lr:
                        layer_code, layer_name = lr
                qty = qty or 1
                pu = price_per_unit or 0
                tp = total_price if total_price is not None else pu * qty
                items_total += tp or 0
                items.append({
                    'layer_name': layer_name,
                    'layer_code': layer_code,
                    'quantity': qty,
                    'price_per_unit': pu,
                    'total_price': tp or 0,
                })

        layers_height = o.get('layers_height') or 0
        total_height = o.get('total_height') or 0
        grand_total = o.get('total_price')
        if grand_total is None:
            grand_total = items_total + cover_price

        return jsonify({
            'order': order,
            'items': items,
            'layers_height': layers_height,
            'total_height': total_height,
            'grand_total': grand_total,
        })
    except Exception as e:
        logger.exception("admin.py: mattress_order_details")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/orders/mattresses/<int:order_id>/status', methods=['PUT'])
def update_mattress_status(order_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    new_status = (data.get('status') or '').strip()
    if new_status not in STATUSES:
        return jsonify({'error': 'Недопустимый статус'}), 400
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if not _table_exists(cur, 'mattress_orders'):
            return jsonify({'error': 'Заказ не найден'}), 404
        cur.execute(
            "UPDATE mattress_orders SET status = %s WHERE id = %s",
            (new_status, order_id),
        )
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Заказ не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_mattress_status")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/orders/mattresses/<int:order_id>', methods=['DELETE'])
def delete_mattress_order(order_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if not _table_exists(cur, 'mattress_orders'):
            return jsonify({'error': 'Заказ не найден'}), 404
        cur.execute("DELETE FROM mattress_orders WHERE id = %s", (order_id,))
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Заказ не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_mattress_order")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#              ДРУГИЕ ЗАКАЗЫ КЛИЕНТА
# ============================================================
@admin_bp.route('/api/admin/orders/customer/<phone>')
def customer_orders(phone):
    if not _require_admin():
        return _unauthorized()
    exclude_type = request.args.get('exclude_type')
    exclude_id = request.args.get('exclude_id', type=int)

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id, user_name, phone, order_date, status FROM orders "
            "WHERE phone = %s ORDER BY order_date DESC",
            (phone,),
        )
        products = []
        for r in cur.fetchall():
            if exclude_type == 'product' and r[0] == exclude_id:
                continue
            products.append({
                'id': r[0],
                'user_name': r[1],
                'phone': r[2],
                'order_date': _format_dt(r[3]),
                'status': r[4],
            })

        mattresses = []
        if _table_exists(cur, 'mattress_orders'):
            cols = _table_columns(cur, 'mattress_orders')
            size_join = ""
            size_select = "'—' AS size_name"
            if 'size_id' in cols and _table_exists(cur, 'mattress_sizes'):
                size_join = "LEFT JOIN mattress_sizes s ON o.size_id = s.id"
                size_select = "COALESCE(s.name, '—') AS size_name"
            date_col = 'order_date' if 'order_date' in cols else 'id'

            cur.execute(f"""
                SELECT o.id, o.{date_col}, o.status, {size_select}
                FROM mattress_orders o
                {size_join}
                WHERE o.phone = %s
                ORDER BY o.{date_col} DESC
            """, (phone,))
            for r in cur.fetchall():
                if exclude_type == 'mattress' and r[0] == exclude_id:
                    continue
                mattresses.append({
                    'id': r[0],
                    'order_date': _format_dt(r[1]),
                    'status': r[2],
                    'size_name': r[3],
                })

        return jsonify({
            'phone': phone,
            'products': products,
            'mattresses': mattresses,
        })
    except Exception as e:
        logger.exception("admin.py: customer_orders")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#                    КЛИЕНТЫ
# ============================================================
@admin_bp.route('/api/admin/clients')
def list_clients():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM users")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + CLIENTS_PAGE_SIZE - 1) // CLIENTS_PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * CLIENTS_PAGE_SIZE

        cur.execute(
            "SELECT id, phone, full_name FROM users ORDER BY id OFFSET %s LIMIT %s",
            (offset, CLIENTS_PAGE_SIZE),
        )
        clients = [{'id': r[0], 'phone': r[1], 'full_name': r[2]} for r in cur.fetchall()]

        return jsonify({
            'clients': clients,
            'total': total,
            'page': page,
            'page_size': CLIENTS_PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_clients")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/clients/search')
def search_clients():
    if not _require_admin():
        return _unauthorized()
    phone = (request.args.get('phone') or '').strip()
    if not phone:
        return jsonify({'error': 'Введите номер телефона'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id, phone, full_name FROM users WHERE phone ILIKE %s ORDER BY id",
            (f'%{phone}%',),
        )
        clients = [{'id': r[0], 'phone': r[1], 'full_name': r[2]} for r in cur.fetchall()]
        return jsonify({'clients': clients})
    except Exception as e:
        logger.exception("admin.py: search_clients")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/clients/<int:client_id>')
def get_client_details(client_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id, phone, full_name, email, address, contact_time, created_at, password_hash "
            "FROM users WHERE id = %s",
            (client_id,),
        )
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Клиент не найден'}), 404

        password_hash = r[7]
        if isinstance(password_hash, (bytes, bytearray)):
            try:
                password_hash = password_hash.decode('utf-8', errors='replace')
            except Exception:
                password_hash = str(password_hash)

        return jsonify({
            'client': {
                'id': r[0],
                'phone': r[1],
                'full_name': r[2],
                'email': r[3],
                'address': r[4],
                'contact_time': r[5],
                'created_at': _format_dt_ru(r[6]),
                'password_hash': password_hash,
            }
        })
    except Exception as e:
        logger.exception("admin.py: get_client_details")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/clients/<int:client_id>', methods=['PUT'])
def update_client(client_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()

    allowed = {'phone', 'full_name', 'email', 'address', 'contact_time', 'password'}
    updates = {}
    for k, v in data.items():
        if k not in allowed:
            continue
        if k == 'password':
            if v:
                updates['password_hash'] = generate_password_hash(v)
        else:
            updates[k] = v

    if not updates:
        return jsonify({'error': 'Нет данных для обновления'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if 'phone' in updates:
            cur.execute("SELECT id FROM users WHERE phone = %s AND id != %s",
                        (updates['phone'], client_id))
            if cur.fetchone():
                return jsonify({'error': 'Такой номер телефона уже занят другим клиентом'}), 400

        set_parts = []
        values = []
        for col, val in updates.items():
            set_parts.append(f"{col} = %s")
            values.append(val)
        values.append(client_id)

        cur.execute(f"UPDATE users SET {', '.join(set_parts)} WHERE id = %s", values)
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Клиент не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_client")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/clients/<int:client_id>', methods=['DELETE'])
def delete_client(client_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM users WHERE id = %s", (client_id,))
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Клиент не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_client")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/clients/<int:client_id>/orders')
def get_client_orders(client_id):
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM orders WHERE user_id = %s", (client_id,))
        total = cur.fetchone()[0]
        total_pages = max(1, (total + CLIENT_ORDERS_PAGE_SIZE - 1) // CLIENT_ORDERS_PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * CLIENT_ORDERS_PAGE_SIZE

        cur.execute(
            """SELECT o.id, o.order_date, o.status,
                      COALESCE((SELECT SUM(cost * quantity) FROM order_items WHERE order_id = o.id), 0) AS total
               FROM orders o
               WHERE o.user_id = %s
               ORDER BY o.order_date DESC
               OFFSET %s LIMIT %s""",
            (client_id, offset, CLIENT_ORDERS_PAGE_SIZE),
        )
        orders = [{
            'id': r[0],
            'order_date': _format_dt_ru(r[1]),
            'status': r[2],
            'total': r[3],
        } for r in cur.fetchall()]

        return jsonify({
            'orders': orders,
            'total': total,
            'page': page,
            'page_size': CLIENT_ORDERS_PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: get_client_orders")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/clients/<int:client_id>/orders/<int:order_id>')
def get_client_order_details(client_id, order_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT id, order_date, status, user_name, phone, email, address,
                   comment, contact_time
            FROM orders WHERE id = %s AND user_id = %s
        """, (order_id, client_id))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Заказ не найден'}), 404

        order = {
            'id': r[0],
            'order_date': _format_dt_ru(r[1]),
            'status': r[2],
            'user_name': r[3],
            'phone': r[4],
            'email': r[5],
            'address': r[6],
            'comment': r[7],
            'contact_time': r[8],
        }

        cur.execute("""
            SELECT product_code, product_name, material_code, material_name,
                   cost, quantity
            FROM order_items WHERE order_id = %s ORDER BY id
        """, (order_id,))
        items = []
        total_sum = 0
        for it in cur.fetchall():
            qty = it[5] or 1
            item_total = (it[4] or 0) * qty
            total_sum += item_total
            items.append({
                'product_code': it[0],
                'product_name': it[1],
                'material_code': it[2],
                'material_name': it[3],
                'cost': it[4] or 0,
                'quantity': qty,
                'total': item_total,
            })

        return jsonify({
            'order': order,
            'items': items,
            'total_sum': total_sum,
        })
    except Exception as e:
        logger.exception("admin.py: get_client_order_details")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#                    ГАЛЕРЕЯ
# ============================================================
def _gallery_section_info(section):
    return GALLERY_SECTIONS.get(section)


def _gallery_validate_name(name):
    if not name:
        return False
    if '/' in name or '\\' in name or '..' in name:
        return False
    return True


def _gallery_projects(section):
    info = _gallery_section_info(section)
    if not info:
        return []
    base = info['base']
    if not base.exists():
        return []
    items = [p.name for p in base.iterdir() if p.is_dir()]
    items.sort(reverse=True)
    return items


def _gallery_project_path(section, name):
    info = _gallery_section_info(section)
    if not info:
        return None
    return info['base'] / name


def _gallery_photo_paths(section, name):
    d = _gallery_project_path(section, name)
    if not d or not d.exists():
        return []
    files = [p for p in d.glob("*.webp") if p.stem.isdigit()]
    files.sort(key=lambda p: int(p.stem))
    return files


def _gallery_read_description(section, name):
    d = _gallery_project_path(section, name)
    if not d:
        return ""
    path = d / "description.txt"
    if path.exists():
        try:
            return path.read_text(encoding="utf-8").strip()
        except Exception:
            return ""
    return ""


def _gallery_write_description(section, name, text):
    d = _gallery_project_path(section, name)
    if not d:
        return
    d.mkdir(parents=True, exist_ok=True)
    _ensure_write_permission(d)
    path = d / "description.txt"
    path.write_text(text, encoding="utf-8")
    _ensure_write_permission(path)


def _gallery_read_captions(section, name):
    d = _gallery_project_path(section, name)
    if not d:
        return {}
    path = d / "captions.json"
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def _gallery_write_caption(section, name, index, caption):
    d = _gallery_project_path(section, name)
    if not d:
        return
    d.mkdir(parents=True, exist_ok=True)
    _ensure_write_permission(d)
    captions = _gallery_read_captions(section, name)
    if caption:
        captions[str(index)] = caption
    else:
        captions.pop(str(index), None)
    path = d / "captions.json"
    path.write_text(
        json.dumps(captions, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    _ensure_write_permission(path)


def _gallery_unique_name(section):
    base = datetime.now().strftime("%Y-%m-%d-%H-%M-%S")
    name = base
    counter = 1
    while _gallery_project_path(section, name).exists():
        name = f"{base}-{counter}"
        counter += 1
        if counter > 1000:
            raise RuntimeError("Не удалось сгенерировать уникальное имя проекта")
    return name


def _ensure_write_permission(path):
    try:
        if path.is_dir():
            os.chmod(path, 0o777)
        else:
            os.chmod(path, 0o777)
    except Exception as e:
        logger.warning(f"Не удалось изменить права для {path}: {e}")


def _remove_readonly(func, path, excinfo):
    try:
        os.chmod(path, 0o777)
        func(path)
    except Exception as e:
        logger.error(f"Не удалось удалить {path}: {e}")
        raise


def _force_delete_directory(path):
    if not path.exists():
        return
    for root, dirs, files in os.walk(path, topdown=False):
        for name in files:
            file_path = os.path.join(root, name)
            try:
                os.chmod(file_path, 0o777)
            except Exception as e:
                logger.warning(f"Не удалось изменить права для файла {file_path}: {e}")
        for name in dirs:
            dir_path = os.path.join(root, name)
            try:
                os.chmod(dir_path, 0o777)
            except Exception as e:
                logger.warning(f"Не удалось изменить права для папки {dir_path}: {e}")
    try:
        os.chmod(path, 0o777)
    except Exception as e:
        logger.warning(f"Не удалось изменить права для корневой папки {path}: {e}")
    shutil.rmtree(path, onerror=_remove_readonly)


def _gallery_save_photo_file(section, name, file_storage, caption=""):
    if not PIL_AVAILABLE:
        logger.error("admin.py: PIL не установлен, невозможно сохранить фото")
        return None
    try:
        project_path = _gallery_project_path(section, name)
        project_path.mkdir(parents=True, exist_ok=True)
        _ensure_write_permission(project_path)

        file_bytes = file_storage.read()
        img = Image.open(io.BytesIO(file_bytes))
        if img.mode in ('RGBA', 'LA', 'P'):
            rgb_img = Image.new('RGB', img.size, (255, 255, 255))
            rgb_img.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else None)
            img = rgb_img

        existing = _gallery_photo_paths(section, name)
        next_num = (max(int(p.stem) for p in existing) + 1) if existing else 1

        filepath = project_path / f"{next_num}.webp"
        img.save(filepath, 'WEBP', quality=85)
        _ensure_write_permission(filepath)

        if caption:
            _gallery_write_caption(section, name, next_num, caption)

        return next_num
    except Exception as e:
        logger.error(f"admin.py: ошибка сохранения фото галереи: {e}")
        return None


@admin_bp.route('/api/admin/gallery/sections')
def gallery_list_sections():
    if not _require_admin():
        return _unauthorized()
    return jsonify({
        'sections': [
            {'key': 'gallery',    'title': GALLERY_SECTIONS['gallery']['title']},
            {'key': 'commercial', 'title': GALLERY_SECTIONS['commercial']['title']},
        ]
    })


@admin_bp.route('/api/admin/gallery/<section>/projects')
def gallery_list_projects(section):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400

    page = max(0, request.args.get('page', 0, type=int))
    projects = _gallery_projects(section)
    total = len(projects)
    total_pages = max(1, (total + GALLERY_PAGE_SIZE - 1) // GALLERY_PAGE_SIZE)
    if page >= total_pages:
        page = total_pages - 1
    offset = page * GALLERY_PAGE_SIZE
    page_items = projects[offset:offset + GALLERY_PAGE_SIZE]

    result = []
    for pname in page_items:
        photos = _gallery_photo_paths(section, pname)
        desc = _gallery_read_description(section, pname)
        result.append({
            'name': pname,
            'description': desc,
            'photo_count': len(photos),
        })

    return jsonify({
        'section': section,
        'title': _gallery_section_info(section)['title'],
        'projects': result,
        'total': total,
        'page': page,
        'page_size': GALLERY_PAGE_SIZE,
        'total_pages': total_pages,
    })


@admin_bp.route('/api/admin/gallery/<section>/projects', methods=['POST'])
def gallery_create_project(section):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not PIL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует библиотека Pillow'}), 500

    try:
        name = _gallery_unique_name(section)
        project_path = _gallery_project_path(section, name)
        project_path.mkdir(parents=True, exist_ok=True)
        _ensure_write_permission(project_path)

        description = (request.form.get('description') or '').strip()
        if description:
            _gallery_write_description(section, name, description)

        files = request.files.getlist('photos')
        saved = []
        for f in files:
            if f and f.filename:
                num = _gallery_save_photo_file(section, name, f)
                if num is not None:
                    saved.append(num)

        return jsonify({
            'success': True,
            'name': name,
            'description': description,
            'photos_saved': saved,
        })
    except Exception as e:
        logger.exception("admin.py: gallery_create_project")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/gallery/<section>/projects/<name>')
def gallery_get_project(section, name):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not _gallery_validate_name(name):
        return jsonify({'error': 'Неверное имя проекта'}), 400
    project_path = _gallery_project_path(section, name)
    if not project_path.exists():
        return jsonify({'error': 'Проект не найден'}), 404

    photos = _gallery_photo_paths(section, name)
    captions = _gallery_read_captions(section, name)
    desc = _gallery_read_description(section, name)

    photos_data = []
    for p in photos:
        idx = int(p.stem)
        photos_data.append({
            'index': idx,
            'url': f"/media/{section}/{name}/{p.stem}.webp",
            'caption': captions.get(p.stem, ''),
        })

    return jsonify({
        'project': {
            'name': name,
            'description': desc,
            'photos': photos_data,
        }
    })


@admin_bp.route('/api/admin/gallery/<section>/projects/<name>/description', methods=['PUT'])
def gallery_update_description(section, name):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not _gallery_validate_name(name):
        return jsonify({'error': 'Неверное имя проекта'}), 400
    project_path = _gallery_project_path(section, name)
    if not project_path.exists():
        return jsonify({'error': 'Проект не найден'}), 404

    data = _json_body()
    text = (data.get('description') or '').strip()
    _gallery_write_description(section, name, text)
    return jsonify({'success': True, 'description': text})


@admin_bp.route('/api/admin/gallery/<section>/projects/<name>/photos', methods=['POST'])
def gallery_upload_photos(section, name):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not _gallery_validate_name(name):
        return jsonify({'error': 'Неверное имя проекта'}), 400
    if not PIL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует библиотека Pillow'}), 500
    project_path = _gallery_project_path(section, name)
    if not project_path.exists():
        return jsonify({'error': 'Проект не найден'}), 404

    files = request.files.getlist('photos')
    if not files:
        return jsonify({'error': 'Не переданы файлы'}), 400

    saved = []
    for f in files:
        if f and f.filename:
            num = _gallery_save_photo_file(section, name, f)
            if num is not None:
                saved.append(num)

    if not saved:
        return jsonify({'error': 'Не удалось сохранить ни одного фото'}), 500
    return jsonify({'success': True, 'saved': saved})


@admin_bp.route('/api/admin/gallery/<section>/projects/<name>/photos/<int:index>/caption', methods=['PUT'])
def gallery_update_caption(section, name, index):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not _gallery_validate_name(name):
        return jsonify({'error': 'Неверное имя проекта'}), 400
    project_path = _gallery_project_path(section, name)
    if not project_path.exists():
        return jsonify({'error': 'Проект не найден'}), 404
    filepath = project_path / f"{index}.webp"
    if not filepath.exists():
        return jsonify({'error': 'Фото не найдено'}), 404

    data = _json_body()
    caption = (data.get('caption') or '').strip()
    _gallery_write_caption(section, name, index, caption)
    return jsonify({'success': True})


@admin_bp.route('/api/admin/gallery/<section>/projects/<name>/photos/<int:index>', methods=['DELETE'])
def gallery_delete_photo(section, name, index):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not _gallery_validate_name(name):
        return jsonify({'error': 'Неверное имя проекта'}), 400
    project_path = _gallery_project_path(section, name)
    if not project_path.exists():
        return jsonify({'error': 'Проект не найден'}), 404
    filepath = project_path / f"{index}.webp"
    if not filepath.exists():
        return jsonify({'error': 'Фото не найдено'}), 404

    try:
        _ensure_write_permission(project_path)
        _ensure_write_permission(filepath)

        filepath.unlink()
        captions = _gallery_read_captions(section, name)
        captions.pop(str(index), None)

        remaining = _gallery_photo_paths(section, name)
        new_caps = {}
        new_idx = 1
        for p in remaining:
            cur = int(p.stem)
            cap = captions.get(str(cur), '')
            if cur != new_idx:
                new_path = project_path / f"{new_idx}.webp"
                _ensure_write_permission(p)
                p.rename(new_path)
            if cap:
                new_caps[str(new_idx)] = cap
            new_idx += 1

        captions_path = project_path / "captions.json"
        if new_caps:
            captions_path.write_text(
                json.dumps(new_caps, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            _ensure_write_permission(captions_path)
        elif captions_path.exists():
            _ensure_write_permission(captions_path)
            captions_path.unlink()

        return jsonify({'success': True})
    except Exception as e:
        logger.exception("admin.py: gallery_delete_photo")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/gallery/<section>/projects/<name>', methods=['DELETE'])
def gallery_delete_project(section, name):
    if not _require_admin():
        return _unauthorized()
    if not _gallery_section_info(section):
        return jsonify({'error': 'Неверный раздел'}), 400
    if not _gallery_validate_name(name):
        return jsonify({'error': 'Неверное имя проекта'}), 400
    project_path = _gallery_project_path(section, name)
    if not project_path.exists():
        return jsonify({'error': 'Проект не найден'}), 404
    try:
        _force_delete_directory(project_path)
        return jsonify({'success': True})
    except Exception as e:
        logger.exception("admin.py: gallery_delete_project")
        return jsonify({'error': str(e)}), 500


# ============================================================
#          МАТРАСЫ: ОБЩИЕ ХЕЛПЕРЫ ДЛЯ ФОТО
# ============================================================
def _layer_photo_path(code):
    return LAYERS_MEDIA / f"{code}.webp"


def _cover_photo_path(code):
    return COVERS_MEDIA / f"{code}.webp"


def _has_layer_photo(code):
    try:
        return _layer_photo_path(code).exists()
    except Exception:
        return False


def _has_cover_photo(code):
    try:
        return _cover_photo_path(code).exists()
    except Exception:
        return False


def _save_image_as_webp(file_storage, target_path):
    """Сохраняет загруженный файл как WebP."""
    if not PIL_AVAILABLE:
        logger.error("admin.py: PIL не установлен")
        return False
    try:
        file_bytes = file_storage.read()
        img = Image.open(io.BytesIO(file_bytes))
        if img.mode in ('RGBA', 'LA', 'P'):
            rgb_img = Image.new('RGB', img.size, (255, 255, 255))
            rgb_img.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else None)
            img = rgb_img
        target_path.parent.mkdir(parents=True, exist_ok=True)
        img.save(target_path, 'WEBP', quality=85)
        _ensure_write_permission(target_path)
        return True
    except Exception as e:
        logger.error(f"admin.py: ошибка сохранения фото: {e}")
        return False


def _delete_file_if_exists(path):
    try:
        if path.exists():
            _ensure_write_permission(path)
            path.unlink()
    except Exception as e:
        logger.warning(f"Не удалось удалить {path}: {e}")


def _move_photo(old_path, new_path):
    try:
        if old_path.exists():
            if new_path.exists():
                _delete_file_if_exists(new_path)
            new_path.parent.mkdir(parents=True, exist_ok=True)
            old_path.rename(new_path)
    except Exception as e:
        logger.warning(f"Не удалось переименовать фото {old_path} → {new_path}: {e}")


# ============================================================
#          МАТРАСЫ: СЛОИ
# ============================================================
@admin_bp.route('/api/admin/mattresses/layers')
def list_layers():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM mattress_layers")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + MATTRESS_PAGE_SIZE - 1) // MATTRESS_PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * MATTRESS_PAGE_SIZE

        cur.execute(
            "SELECT id, code, name, description FROM mattress_layers "
            "ORDER BY code OFFSET %s LIMIT %s",
            (offset, MATTRESS_PAGE_SIZE),
        )
        layers = []
        for r in cur.fetchall():
            layers.append({
                'id': r[0],
                'code': r[1],
                'name': r[2],
                'description': r[3] or '',
                'has_photo': _has_layer_photo(r[1]),
            })

        return jsonify({
            'layers': layers,
            'total': total,
            'page': page,
            'page_size': MATTRESS_PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_layers")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers', methods=['POST'])
def create_layer():
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    code = (data.get('code') or '').strip()
    name = (data.get('name') or '').strip()
    description = (data.get('description') or '').strip()
    if not code or not name:
        return jsonify({'error': 'Код и название обязательны'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT id FROM mattress_layers WHERE code = %s", (code,))
        if cur.fetchone():
            return jsonify({'error': 'Слой с таким кодом уже существует'}), 400
        cur.execute(
            "INSERT INTO mattress_layers (code, name, description) "
            "VALUES (%s, %s, %s) RETURNING id",
            (code, name, description),
        )
        new_id = cur.fetchone()[0]
        conn.commit()
        return jsonify({'success': True, 'id': new_id})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: create_layer")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>')
def get_layer(layer_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id, code, name, description FROM mattress_layers WHERE id = %s",
            (layer_id,),
        )
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Слой не найден'}), 404
        has_photo = _has_layer_photo(r[1])
        return jsonify({
            'layer': {
                'id': r[0],
                'code': r[1],
                'name': r[2],
                'description': r[3] or '',
                'has_photo': has_photo,
                'photo_url': f"/media/layers/{r[1]}.webp" if has_photo else None,
            }
        })
    except Exception as e:
        logger.exception("admin.py: get_layer")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>', methods=['PUT'])
def update_layer(layer_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    code = (data.get('code') or '').strip()
    name = (data.get('name') or '').strip()
    description = (data.get('description') or '').strip()
    if not code or not name:
        return jsonify({'error': 'Код и название обязательны'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_layers WHERE id = %s", (layer_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Слой не найден'}), 404
        old_code = r[0]

        if code != old_code:
            cur.execute(
                "SELECT id FROM mattress_layers WHERE code = %s AND id != %s",
                (code, layer_id),
            )
            if cur.fetchone():
                return jsonify({'error': 'Слой с таким кодом уже существует'}), 400

        cur.execute(
            "UPDATE mattress_layers SET code=%s, name=%s, description=%s WHERE id=%s",
            (code, name, description, layer_id),
        )
        conn.commit()

        if code != old_code:
            _move_photo(_layer_photo_path(old_code), _layer_photo_path(code))

        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_layer")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>', methods=['DELETE'])
def delete_layer(layer_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_layers WHERE id = %s", (layer_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Слой не найден'}), 404
        code = r[0]

        if _table_exists(cur, 'mattress_layers_prices'):
            cur.execute("DELETE FROM mattress_layers_prices WHERE layer_id = %s", (layer_id,))
        if _table_exists(cur, 'mattress_layer_main_features'):
            cur.execute("DELETE FROM mattress_layer_main_features WHERE layer_id = %s", (layer_id,))
        if _table_exists(cur, 'mattress_layer_extra_features'):
            cur.execute("DELETE FROM mattress_layer_extra_features WHERE layer_id = %s", (layer_id,))
        cur.execute("DELETE FROM mattress_layers WHERE id = %s", (layer_id,))
        conn.commit()

        _delete_file_if_exists(_layer_photo_path(code))
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_layer")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>/photo', methods=['POST'])
def upload_layer_photo(layer_id):
    if not _require_admin():
        return _unauthorized()
    if not PIL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует Pillow'}), 500

    file = request.files.get('photo')
    if not file or not file.filename:
        return jsonify({'error': 'Не передан файл'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_layers WHERE id = %s", (layer_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Слой не найден'}), 404
        code = r[0]
        target = _layer_photo_path(code)
        if not _save_image_as_webp(file, target):
            return jsonify({'error': 'Не удалось сохранить фото'}), 500
        return jsonify({'success': True, 'photo_url': f"/media/layers/{code}.webp"})
    except Exception as e:
        logger.exception("admin.py: upload_layer_photo")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>/photo', methods=['DELETE'])
def remove_layer_photo(layer_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_layers WHERE id = %s", (layer_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Слой не найден'}), 404
        _delete_file_if_exists(_layer_photo_path(r[0]))
        return jsonify({'success': True})
    except Exception as e:
        logger.exception("admin.py: remove_layer_photo")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>/main-features')
def list_layer_main_features(layer_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT mft.id as feature_id, mft.name as feature_name,
                   COALESCE(pmf.value, '') as value
            FROM mattress_layer_main_features_types mft
            LEFT JOIN mattress_layer_main_features pmf
                ON mft.id = pmf.feature_id AND pmf.layer_id = %s
            ORDER BY mft.name
        """, (layer_id,))
        features = [{
            'feature_id': r[0],
            'feature_name': r[1],
            'value': r[2] or '',
        } for r in cur.fetchall()]
        return jsonify({'features': features})
    except Exception as e:
        logger.exception("admin.py: list_layer_main_features")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>/main-features/<int:feature_id>', methods=['PUT'])
def set_layer_main_feature(layer_id, feature_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    value = (data.get('value') or '').strip()

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT INTO mattress_layer_main_features (layer_id, feature_id, value)
            VALUES (%s, %s, %s)
            ON CONFLICT (layer_id, feature_id) DO UPDATE SET value = EXCLUDED.value
        """, (layer_id, feature_id, value))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: set_layer_main_feature")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>/extra-features')
def list_layer_extra_features(layer_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id, name, value FROM mattress_layer_extra_features "
            "WHERE layer_id = %s ORDER BY name",
            (layer_id,),
        )
        features = [{'id': r[0], 'name': r[1], 'value': r[2] or ''} for r in cur.fetchall()]
        return jsonify({'features': features})
    except Exception as e:
        logger.exception("admin.py: list_layer_extra_features")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/layers/<int:layer_id>/extra-features', methods=['POST'])
def add_layer_extra_feature(layer_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    name = (data.get('name') or '').strip()
    value = (data.get('value') or '').strip()
    if not name:
        return jsonify({'error': 'Название обязательно'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT id FROM mattress_layers WHERE id = %s", (layer_id,))
        if not cur.fetchone():
            return jsonify({'error': 'Слой не найден'}), 404
        cur.execute(
            "INSERT INTO mattress_layer_extra_features (layer_id, name, value) "
            "VALUES (%s, %s, %s) RETURNING id",
            (layer_id, name, value),
        )
        new_id = cur.fetchone()[0]
        conn.commit()
        return jsonify({'success': True, 'id': new_id})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: add_layer_extra_feature")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/extra-features/<int:feat_id>', methods=['PUT'])
def update_extra_feature(feat_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    updates = {}
    if 'name' in data:
        updates['name'] = (data.get('name') or '').strip()
    if 'value' in data:
        updates['value'] = (data.get('value') or '').strip()
    if not updates:
        return jsonify({'error': 'Нет данных'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        set_parts = []
        values = []
        for col, val in updates.items():
            set_parts.append(f"{col} = %s")
            values.append(val)
        values.append(feat_id)
        cur.execute(
            f"UPDATE mattress_layer_extra_features SET {', '.join(set_parts)} WHERE id = %s",
            values,
        )
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Характеристика не найдена'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_extra_feature")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/extra-features/<int:feat_id>', methods=['DELETE'])
def delete_extra_feature(feat_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM mattress_layer_extra_features WHERE id = %s", (feat_id,))
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Характеристика не найдена'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_extra_feature")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#          МАТРАСЫ: ЧЕХЛЫ
# ============================================================
@admin_bp.route('/api/admin/mattresses/covers')
def list_covers():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM mattress_cover")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + MATTRESS_PAGE_SIZE - 1) // MATTRESS_PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * MATTRESS_PAGE_SIZE

        cur.execute(
            "SELECT id, code, name, description FROM mattress_cover "
            "ORDER BY code OFFSET %s LIMIT %s",
            (offset, MATTRESS_PAGE_SIZE),
        )
        covers = []
        for r in cur.fetchall():
            covers.append({
                'id': r[0],
                'code': r[1],
                'name': r[2],
                'description': r[3] or '',
                'has_photo': _has_cover_photo(r[1]),
            })

        return jsonify({
            'covers': covers,
            'total': total,
            'page': page,
            'page_size': MATTRESS_PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_covers")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/covers', methods=['POST'])
def create_cover():
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    code = (data.get('code') or '').strip()
    name = (data.get('name') or '').strip()
    description = (data.get('description') or '').strip()
    if not code or not name:
        return jsonify({'error': 'Код и название обязательны'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT id FROM mattress_cover WHERE code = %s", (code,))
        if cur.fetchone():
            return jsonify({'error': 'Чехол с таким кодом уже существует'}), 400
        cur.execute(
            "INSERT INTO mattress_cover (code, name, description) "
            "VALUES (%s, %s, %s) RETURNING id",
            (code, name, description),
        )
        new_id = cur.fetchone()[0]
        conn.commit()
        return jsonify({'success': True, 'id': new_id})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: create_cover")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/covers/<int:cover_id>')
def get_cover(cover_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id, code, name, description FROM mattress_cover WHERE id = %s",
            (cover_id,),
        )
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Чехол не найден'}), 404
        has_photo = _has_cover_photo(r[1])
        return jsonify({
            'cover': {
                'id': r[0],
                'code': r[1],
                'name': r[2],
                'description': r[3] or '',
                'has_photo': has_photo,
                'photo_url': f"/media/covers/{r[1]}.webp" if has_photo else None,
            }
        })
    except Exception as e:
        logger.exception("admin.py: get_cover")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/covers/<int:cover_id>', methods=['PUT'])
def update_cover(cover_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    code = (data.get('code') or '').strip()
    name = (data.get('name') or '').strip()
    description = (data.get('description') or '').strip()
    if not code or not name:
        return jsonify({'error': 'Код и название обязательны'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_cover WHERE id = %s", (cover_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Чехол не найден'}), 404
        old_code = r[0]

        if code != old_code:
            cur.execute(
                "SELECT id FROM mattress_cover WHERE code = %s AND id != %s",
                (code, cover_id),
            )
            if cur.fetchone():
                return jsonify({'error': 'Чехол с таким кодом уже существует'}), 400

        cur.execute(
            "UPDATE mattress_cover SET code=%s, name=%s, description=%s WHERE id=%s",
            (code, name, description, cover_id),
        )
        conn.commit()

        if code != old_code:
            _move_photo(_cover_photo_path(old_code), _cover_photo_path(code))

        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_cover")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/covers/<int:cover_id>', methods=['DELETE'])
def delete_cover(cover_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_cover WHERE id = %s", (cover_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Чехол не найден'}), 404
        code = r[0]

        if _table_exists(cur, 'mattress_cover_prices'):
            cur.execute("DELETE FROM mattress_cover_prices WHERE cover_id = %s", (cover_id,))
        cur.execute("DELETE FROM mattress_cover WHERE id = %s", (cover_id,))
        conn.commit()

        _delete_file_if_exists(_cover_photo_path(code))
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_cover")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/covers/<int:cover_id>/photo', methods=['POST'])
def upload_cover_photo(cover_id):
    if not _require_admin():
        return _unauthorized()
    if not PIL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует Pillow'}), 500

    file = request.files.get('photo')
    if not file or not file.filename:
        return jsonify({'error': 'Не передан файл'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_cover WHERE id = %s", (cover_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Чехол не найден'}), 404
        code = r[0]
        target = _cover_photo_path(code)
        if not _save_image_as_webp(file, target):
            return jsonify({'error': 'Не удалось сохранить фото'}), 500
        return jsonify({'success': True, 'photo_url': f"/media/covers/{code}.webp"})
    except Exception as e:
        logger.exception("admin.py: upload_cover_photo")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/covers/<int:cover_id>/photo', methods=['DELETE'])
def remove_cover_photo(cover_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT code FROM mattress_cover WHERE id = %s", (cover_id,))
        r = cur.fetchone()
        if not r:
            return jsonify({'error': 'Чехол не найден'}), 404
        _delete_file_if_exists(_cover_photo_path(r[0]))
        return jsonify({'success': True})
    except Exception as e:
        logger.exception("admin.py: remove_cover_photo")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#          МАТРАСЫ: РАЗМЕРЫ
# ============================================================
@admin_bp.route('/api/admin/mattresses/sizes')
def list_sizes():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM mattress_sizes")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + MATTRESS_PAGE_SIZE - 1) // MATTRESS_PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * MATTRESS_PAGE_SIZE

        cur.execute(
            "SELECT id, name FROM mattress_sizes ORDER BY name OFFSET %s LIMIT %s",
            (offset, MATTRESS_PAGE_SIZE),
        )
        sizes = [{'id': r[0], 'name': r[1]} for r in cur.fetchall()]

        return jsonify({
            'sizes': sizes,
            'total': total,
            'page': page,
            'page_size': MATTRESS_PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_sizes")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/sizes', methods=['POST'])
def create_size():
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    name = (data.get('name') or '').strip()
    if not name:
        return jsonify({'error': 'Название обязательно'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT id FROM mattress_sizes WHERE name = %s", (name,))
        if cur.fetchone():
            return jsonify({'error': 'Размер с таким названием уже существует'}), 400
        cur.execute(
            "INSERT INTO mattress_sizes (name) VALUES (%s) RETURNING id",
            (name,),
        )
        new_id = cur.fetchone()[0]
        conn.commit()
        return jsonify({'success': True, 'id': new_id})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: create_size")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/sizes/<int:size_id>', methods=['PUT'])
def update_size(size_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    name = (data.get('name') or '').strip()
    if not name:
        return jsonify({'error': 'Название обязательно'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id FROM mattress_sizes WHERE name = %s AND id != %s",
            (name, size_id),
        )
        if cur.fetchone():
            return jsonify({'error': 'Размер с таким названием уже существует'}), 400
        cur.execute(
            "UPDATE mattress_sizes SET name = %s WHERE id = %s",
            (name, size_id),
        )
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Размер не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_size")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/sizes/<int:size_id>', methods=['DELETE'])
def delete_size(size_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if _table_exists(cur, 'mattress_layers_prices'):
            cur.execute("DELETE FROM mattress_layers_prices WHERE size_id = %s", (size_id,))
        if _table_exists(cur, 'mattress_cover_prices'):
            cur.execute("DELETE FROM mattress_cover_prices WHERE size_id = %s", (size_id,))
        cur.execute("DELETE FROM mattress_sizes WHERE id = %s", (size_id,))
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Размер не найден'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_size")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#          МАТРАСЫ: ЦЕНЫ
# ============================================================
def _price_config(price_type):
    if price_type == 'layers':
        return 'mattress_layers_prices', 'layer_id'
    if price_type == 'covers':
        return 'mattress_cover_prices', 'cover_id'
    return None, None


@admin_bp.route('/api/admin/mattresses/prices/<price_type>/<int:entity_id>')
def get_prices(price_type, entity_id):
    if not _require_admin():
        return _unauthorized()
    price_table, entity_col = _price_config(price_type)
    if not price_table:
        return jsonify({'error': 'Неверный тип цен'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(f"""
            SELECT s.id, s.name, p.price
            FROM mattress_sizes s
            LEFT JOIN {price_table} p
                ON p.size_id = s.id AND p.{entity_col} = %s
            ORDER BY s.name
        """, (entity_id,))
        prices = [{
            'size_id': r[0],
            'size_name': r[1],
            'price': r[2],
        } for r in cur.fetchall()]

        return jsonify({
            'prices': prices,
            'price_type': price_type,
            'entity_id': entity_id,
        })
    except Exception as e:
        logger.exception("admin.py: get_prices")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route(
    '/api/admin/mattresses/prices/<price_type>/<int:entity_id>/<int:size_id>',
    methods=['PUT'],
)
def set_price(price_type, entity_id, size_id):
    if not _require_admin():
        return _unauthorized()
    price_table, entity_col = _price_config(price_type)
    if not price_table:
        return jsonify({'error': 'Неверный тип цен'}), 400

    data = _json_body()
    price = data.get('price')
    if price is not None and price != '':
        try:
            price = int(price)
        except (ValueError, TypeError):
            return jsonify({'error': 'Неверная цена'}), 400
        if price < 0:
            return jsonify({'error': 'Цена должна быть неотрицательной'}), 400
    else:
        price = None

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if price is None:
            cur.execute(
                f"DELETE FROM {price_table} WHERE {entity_col} = %s AND size_id = %s",
                (entity_id, size_id),
            )
        else:
            cur.execute(f"""
                INSERT INTO {price_table} ({entity_col}, size_id, price)
                VALUES (%s, %s, %s)
                ON CONFLICT ({entity_col}, size_id)
                DO UPDATE SET price = EXCLUDED.price
            """, (entity_id, size_id, price))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: set_price")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#          МАТРАСЫ: СПРАВОЧНИК ХАРАКТЕРИСТИК
# ============================================================
@admin_bp.route('/api/admin/mattresses/feature-types')
def list_feature_types():
    if not _require_admin():
        return _unauthorized()
    page = max(0, request.args.get('page', 0, type=int))
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM mattress_layer_main_features_types")
        total = cur.fetchone()[0]
        total_pages = max(1, (total + MATTRESS_PAGE_SIZE - 1) // MATTRESS_PAGE_SIZE)
        if page >= total_pages:
            page = total_pages - 1
        offset = page * MATTRESS_PAGE_SIZE

        cur.execute(
            "SELECT id, name FROM mattress_layer_main_features_types "
            "ORDER BY name OFFSET %s LIMIT %s",
            (offset, MATTRESS_PAGE_SIZE),
        )
        types = [{'id': r[0], 'name': r[1]} for r in cur.fetchall()]

        return jsonify({
            'types': types,
            'total': total,
            'page': page,
            'page_size': MATTRESS_PAGE_SIZE,
            'total_pages': total_pages,
        })
    except Exception as e:
        logger.exception("admin.py: list_feature_types")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/feature-types', methods=['POST'])
def create_feature_type():
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    name = (data.get('name') or '').strip()
    if not name:
        return jsonify({'error': 'Название обязательно'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id FROM mattress_layer_main_features_types WHERE name = %s",
            (name,),
        )
        if cur.fetchone():
            return jsonify({'error': 'Такая характеристика уже существует'}), 400
        cur.execute(
            "INSERT INTO mattress_layer_main_features_types (name) "
            "VALUES (%s) RETURNING id",
            (name,),
        )
        new_id = cur.fetchone()[0]
        conn.commit()
        return jsonify({'success': True, 'id': new_id})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: create_feature_type")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/feature-types/<int:feat_id>', methods=['PUT'])
def update_feature_type(feat_id):
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    name = (data.get('name') or '').strip()
    if not name:
        return jsonify({'error': 'Название обязательно'}), 400

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        cur.execute(
            "SELECT id FROM mattress_layer_main_features_types "
            "WHERE name = %s AND id != %s",
            (name, feat_id),
        )
        if cur.fetchone():
            return jsonify({'error': 'Такая характеристика уже существует'}), 400
        cur.execute(
            "UPDATE mattress_layer_main_features_types SET name = %s WHERE id = %s",
            (name, feat_id),
        )
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Характеристика не найдена'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: update_feature_type")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/mattresses/feature-types/<int:feat_id>', methods=['DELETE'])
def delete_feature_type(feat_id):
    if not _require_admin():
        return _unauthorized()
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        if _table_exists(cur, 'mattress_layer_main_features'):
            cur.execute(
                "DELETE FROM mattress_layer_main_features WHERE feature_id = %s",
                (feat_id,),
            )
        cur.execute(
            "DELETE FROM mattress_layer_main_features_types WHERE id = %s",
            (feat_id,),
        )
        affected = cur.rowcount
        conn.commit()
        if affected == 0:
            return jsonify({'error': 'Характеристика не найдена'}), 404
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        logger.exception("admin.py: delete_feature_type")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#          МАТРАСЫ: ЭКСПОРТ EXCEL (000000.xlsx)
# ============================================================
def _build_mattress_export_workbook(cur):
    """Строит workbook 000000.xlsx из данных БД."""
    cur.execute("SELECT id, name FROM mattress_sizes ORDER BY name")
    sizes = cur.fetchall()
    cur.execute("SELECT id, code FROM mattress_layers ORDER BY code")
    layers = cur.fetchall()
    cur.execute("SELECT id, code FROM mattress_cover ORDER BY code")
    covers = cur.fetchall()

    cur.execute("SELECT layer_id, size_id, price FROM mattress_layers_prices")
    layer_prices = {(r[0], r[1]): r[2] for r in cur.fetchall()}
    cur.execute("SELECT cover_id, size_id, price FROM mattress_cover_prices")
    cover_prices = {(r[0], r[1]): r[2] for r in cur.fetchall()}

    wb = Workbook()
    ws = wb.active
    ws.title = "000000"

    headers = ["Размер"] + [l[1] for l in layers] + [c[1] for c in covers]
    ws.append(headers)

    for size in sizes:
        row = [size[1]]
        for layer in layers:
            row.append(layer_prices.get((layer[0], size[0]), ""))
        for cover in covers:
            row.append(cover_prices.get((cover[0], size[0]), ""))
        ws.append(row)

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = header_fill
        cell.border = border

    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row_cells:
            cell.border = border
            if cell.column == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(min(max_len + 2, 50), 50)

    return wb


@admin_bp.route('/api/admin/mattresses/export')
def export_mattresses_excel():
    if not _require_admin():
        return _unauthorized()
    if not OPENPYXL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует openpyxl'}), 500

    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        wb = _build_mattress_export_workbook(cur)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='000000.xlsx',
        )
    except Exception as e:
        logger.exception("admin.py: export_mattresses_excel")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ============================================================
#          ПАРТНЁРСКИЕ ЦЕНЫ (аналог partners_prices_bot.py)
# ============================================================
_PARTNER_NAME_RE = re.compile(r"[0-9a-zA-Zа-яА-ЯёЁ_\-]+")


def _partner_file_path(name):
    return PARTNERS_BASE / f"{name}.xlsx"


def _partner_is_valid_name(name):
    if not name or len(name) > 40:
        return False
    if name == ORIGINAL_PARTNER_NAME:
        return False
    return bool(_PARTNER_NAME_RE.fullmatch(name))


def _partner_list_files():
    if not PARTNERS_BASE.exists():
        return []
    files = [f for f in PARTNERS_BASE.glob("*.xlsx") if f.is_file()]
    files.sort(key=lambda p: (p.stem != ORIGINAL_PARTNER_NAME, p.stem.lower()))
    return files


def _ensure_partner_template_file():
    """Гарантирует наличие partners/000000.xlsx. Возвращает True при успехе."""
    path = _partner_file_path(ORIGINAL_PARTNER_NAME)
    if path.exists():
        return True
    if not OPENPYXL_AVAILABLE:
        return False
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        try:
            wb = _build_mattress_export_workbook(cur)
            PARTNERS_BASE.mkdir(parents=True, exist_ok=True)
            wb.save(path)
            return True
        finally:
            cur.close()
            conn.close()
    except Exception as e:
        logger.warning(f"admin.py: не удалось создать шаблон 000000.xlsx: {e}")
        return False


def _validate_partner_upload(new_path, template_path):
    """Проверяет структуру загруженного xlsx по шаблону."""
    if not template_path.exists():
        return True, "OK"
    try:
        wb1 = load_workbook(template_path)
        ws1 = wb1.active
        wb2 = load_workbook(new_path)
        ws2 = wb2.active
    except Exception as e:
        return False, f"не удалось прочитать файл ({e})"

    if ws1.max_row != ws2.max_row:
        return False, (f"число строк не совпадает "
                       f"(ожидается {ws1.max_row}, получено {ws2.max_row})")
    if ws1.max_column != ws2.max_column:
        return False, (f"число столбцов не совпадает "
                       f"(ожидается {ws1.max_column}, получено {ws2.max_column})")

    for c in range(2, ws1.max_column + 1):
        v1 = ws1.cell(1, c).value
        v2 = ws2.cell(1, c).value
        if str(v1) != str(v2):
            return False, f"заголовок столбца {c} не совпадает ('{v1}' vs '{v2}')"

    for r in range(2, ws1.max_row + 1):
        v1 = ws1.cell(r, 1).value
        v2 = ws2.cell(r, 1).value
        if str(v1) != str(v2):
            return False, f"метка строки {r} не совпадает ('{v1}' vs '{v2}')"

    for r in range(2, ws1.max_row + 1):
        for c in range(2, ws1.max_column + 1):
            v = ws2.cell(r, c).value
            if v is None or v == "":
                return False, f"ячейка ({r},{c}) пуста"
            try:
                float(v)
            except (ValueError, TypeError):
                return False, f"ячейка ({r},{c}) не является числом"

    return True, "OK"


@admin_bp.route('/api/admin/partners/files')
def partner_list_files():
    if not _require_admin():
        return _unauthorized()
    try:
        files = _partner_list_files()
        result = []
        for f in files:
            try:
                size_kb = round(f.stat().st_size / 1024, 1)
            except Exception:
                size_kb = 0
            result.append({
                'name': f.stem,
                'is_original': f.stem == ORIGINAL_PARTNER_NAME,
                'size_kb': size_kb,
            })
        return jsonify({'files': result})
    except Exception as e:
        logger.exception("admin.py: partner_list_files")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/template')
def partner_download_template():
    """Генерирует 000000.xlsx из БД, сохраняет в partners/ и отдаёт на скачивание."""
    if not _require_admin():
        return _unauthorized()
    if not OPENPYXL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует openpyxl'}), 500
    try:
        conn = _get_db_connection()
    except Exception as e:
        logger.exception("admin.py: ошибка подключения к БД")
        return jsonify({'error': f'Ошибка подключения к БД: {e}'}), 500
    cur = conn.cursor()
    try:
        wb = _build_mattress_export_workbook(cur)
        buffer = io.BytesIO()
        wb.save(buffer)
        content = buffer.getvalue()

        PARTNERS_BASE.mkdir(parents=True, exist_ok=True)
        template_path = _partner_file_path(ORIGINAL_PARTNER_NAME)
        with open(template_path, 'wb') as f:
            f.write(content)

        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='000000.xlsx',
        )
    except Exception as e:
        logger.exception("admin.py: partner_download_template")
        return jsonify({'error': str(e)}), 500
    finally:
        cur.close()
        conn.close()


@admin_bp.route('/api/admin/partners/files', methods=['POST'])
def partner_create_copy():
    if not _require_admin():
        return _unauthorized()
    data = _json_body()
    name = (data.get('name') or '').strip()
    if not _partner_is_valid_name(name):
        return jsonify({'error': 'Недопустимое имя файла'}), 400
    dst = _partner_file_path(name)
    if dst.exists():
        return jsonify({'error': 'Файл с таким именем уже существует'}), 400

    if not _ensure_partner_template_file():
        return jsonify({'error': 'Шаблон 000000.xlsx не найден. Нажмите «Скачать 000000.xlsx».'}), 400

    src = _partner_file_path(ORIGINAL_PARTNER_NAME)
    try:
        shutil.copy(str(src), str(dst))
        return jsonify({'success': True, 'name': name})
    except Exception as e:
        logger.exception("admin.py: partner_create_copy")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/files/<name>/download')
def partner_download_file(name):
    if not _require_admin():
        return _unauthorized()
    path = _partner_file_path(name)
    if not path.exists():
        return jsonify({'error': 'Файл не найден'}), 404
    try:
        return send_file(
            str(path),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'{name}.xlsx',
        )
    except Exception as e:
        logger.exception("admin.py: partner_download_file")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/files/<name>/grid')
def partner_file_grid(name):
    if not _require_admin():
        return _unauthorized()
    path = _partner_file_path(name)
    if not path.exists():
        return jsonify({'error': 'Файл не найден'}), 404
    if not OPENPYXL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует openpyxl'}), 500
    try:
        wb = load_workbook(path)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        row_labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        values = []
        for r in range(1, ws.max_row + 1):
            row_vals = []
            for c in range(1, ws.max_column + 1):
                row_vals.append(ws.cell(r, c).value)
            values.append(row_vals)
        return jsonify({
            'name': name,
            'is_original': name == ORIGINAL_PARTNER_NAME,
            'headers': headers,
            'row_labels': row_labels,
            'values': values,
            'max_row': ws.max_row,
            'max_column': ws.max_column,
        })
    except Exception as e:
        logger.exception("admin.py: partner_file_grid")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/files/<name>/bulk', methods=['PUT'])
def partner_bulk_update(name):
    if not _require_admin():
        return _unauthorized()
    if name == ORIGINAL_PARTNER_NAME:
        return jsonify({'error': 'Нельзя изменять шаблон'}), 400
    path = _partner_file_path(name)
    if not path.exists():
        return jsonify({'error': 'Файл не найден'}), 404
    if not OPENPYXL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует openpyxl'}), 500

    data = _json_body()
    try:
        pct = float(data.get('percent'))
    except (ValueError, TypeError):
        return jsonify({'error': 'Неверный процент'}), 400

    try:
        wb = load_workbook(path)
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            for c in range(2, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is None or v == "":
                    continue
                try:
                    old = float(v)
                except (ValueError, TypeError):
                    continue
                ws.cell(r, c).value = int(round(old * (1 + pct / 100.0)))
        wb.save(path)
        return jsonify({'success': True})
    except Exception as e:
        logger.exception("admin.py: partner_bulk_update")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/files/<name>/cell', methods=['PUT'])
def partner_update_cell(name):
    if not _require_admin():
        return _unauthorized()
    if name == ORIGINAL_PARTNER_NAME:
        return jsonify({'error': 'Нельзя изменять шаблон'}), 400
    path = _partner_file_path(name)
    if not path.exists():
        return jsonify({'error': 'Файл не найден'}), 404
    if not OPENPYXL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует openpyxl'}), 500

    data = _json_body()
    try:
        row = int(data.get('row'))
        col = int(data.get('col'))
        value = float(data.get('value'))
    except (ValueError, TypeError):
        return jsonify({'error': 'Неверные параметры'}), 400
    op = data.get('op')
    if op not in ('set', 'mul', 'add', 'sub'):
        return jsonify({'error': 'Неизвестная операция'}), 400

    try:
        wb = load_workbook(path)
        ws = wb.active
        if row < 2 or row > ws.max_row or col < 2 or col > ws.max_column:
            return jsonify({'error': 'Неверные координаты ячейки'}), 400
        old_raw = ws.cell(row, col).value
        try:
            old = float(old_raw)
        except (ValueError, TypeError):
            old = 0.0

        if op == 'set':
            new = value
        elif op == 'mul':
            new = old * value
        elif op == 'add':
            new = old + value
        else:  # sub
            new = old - value

        new_int = int(round(new))
        ws.cell(row, col).value = new_int
        wb.save(path)
        return jsonify({'success': True, 'old': old, 'new': new_int})
    except Exception as e:
        logger.exception("admin.py: partner_update_cell")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/files/<name>/rename', methods=['PUT'])
def partner_rename_file(name):
    if not _require_admin():
        return _unauthorized()
    if name == ORIGINAL_PARTNER_NAME:
        return jsonify({'error': 'Нельзя переименовать шаблон'}), 400
    data = _json_body()
    new_name = (data.get('new_name') or '').strip()
    if not _partner_is_valid_name(new_name):
        return jsonify({'error': 'Недопустимое имя'}), 400
    src = _partner_file_path(name)
    dst = _partner_file_path(new_name)
    if not src.exists():
        return jsonify({'error': 'Файл не найден'}), 404
    if dst.exists():
        return jsonify({'error': 'Файл с таким именем уже существует'}), 400
    try:
        src.rename(dst)
        return jsonify({'success': True, 'name': new_name})
    except Exception as e:
        logger.exception("admin.py: partner_rename_file")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/files/<name>', methods=['DELETE'])
def partner_delete_file(name):
    if not _require_admin():
        return _unauthorized()
    if name == ORIGINAL_PARTNER_NAME:
        return jsonify({'error': 'Нельзя удалить шаблон 000000.xlsx'}), 400
    path = _partner_file_path(name)
    if not path.exists():
        return jsonify({'error': 'Файл не найден'}), 404
    try:
        path.unlink()
        return jsonify({'success': True})
    except Exception as e:
        logger.exception("admin.py: partner_delete_file")
        return jsonify({'error': str(e)}), 500


@admin_bp.route('/api/admin/partners/upload', methods=['POST'])
def partner_upload():
    if not _require_admin():
        return _unauthorized()
    if not OPENPYXL_AVAILABLE:
        return jsonify({'error': 'На сервере отсутствует openpyxl'}), 500

    file = request.files.get('file')
    if not file or not file.filename:
        return jsonify({'error': 'Файл не передан'}), 400
    if not file.filename.lower().endswith('.xlsx'):
        return jsonify({'error': 'Поддерживаются только .xlsx файлы'}), 400

    mode = (request.form.get('mode') or 'new').strip()
    target_name = (request.form.get('name') or '').strip()

    tmp_dir = Path(tempfile.mkdtemp(prefix='partner_upload_'))
    tmp_path = tmp_dir / file.filename
    try:
        file.save(str(tmp_path))

        if not _ensure_partner_template_file():
            return jsonify({'error': 'Не удалось получить шаблон для проверки'}), 500

        template_path = _partner_file_path(ORIGINAL_PARTNER_NAME)
        ok, msg = _validate_partner_upload(tmp_path, template_path)
        if not ok:
            return jsonify({'error': f'Файл не прошёл проверку: {msg}'}), 400

        if mode == 'replace':
            if not target_name or target_name == ORIGINAL_PARTNER_NAME:
                return jsonify({'error': 'Нельзя заменить шаблон'}), 400
            if not _partner_is_valid_name(target_name):
                return jsonify({'error': 'Недопустимое имя'}), 400
            dst = _partner_file_path(target_name)
            if not dst.exists():
                return jsonify({'error': 'Файл для замены не найден'}), 404
        else:  # mode == 'new'
            if not _partner_is_valid_name(target_name):
                return jsonify({'error': 'Недопустимое имя'}), 400
            dst = _partner_file_path(target_name)
            if dst.exists():
                return jsonify({'error': 'Файл с таким именем уже существует'}), 400

        shutil.copy(str(tmp_path), str(dst))
        return jsonify({'success': True, 'name': target_name})
    except Exception as e:
        logger.exception("admin.py: partner_upload")
        return jsonify({'error': str(e)}), 500
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)