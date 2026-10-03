#!/usr/bin/env python3
"""
Telegram-бот для управления файлами цен партнёров.

- Работает без БД.
- Базируется на шаблоне /partners/000000.xlsx.
- Позволяет:
    * копировать 000000.xlsx под новым именем,
    * менять все цены на процент,
    * менять одну цену (set / mul / add / sub),
    * переименовывать и удалять созданные файлы,
    * загружать свой .xlsx, если структура совпадает с 000000.xlsx
      и все ценовые ячейки заполнены.
- Шаблон 000000.xlsx защищён от изменения/удаления.
"""

import logging
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import List, Optional, Tuple

from openpyxl import load_workbook

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ConversationHandler,
    MessageHandler,
    filters,
    ContextTypes,
)
from telegram.request import HTTPXRequest

# ------------------------- Пути и конфигурация -------------------------
BASE_DIR = Path(__file__).parent
TOKEN_FILE = BASE_DIR / "token_partners_prices.txt"
ADMIN_FILE = BASE_DIR / "admin_ids.txt"

# partners/000000.xlsx лежит на уровень выше telegram_admin_panel/
PARTNERS_DIR = BASE_DIR.parent / "partners"
ORIGINAL_NAME = "000000"
ORIGINAL_FILE = PARTNERS_DIR / f"{ORIGINAL_NAME}.xlsx"

PARTNERS_DIR.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)


def read_token() -> str:
    if not TOKEN_FILE.exists():
        logger.error(f"Файл с токеном не найден: {TOKEN_FILE}")
        sys.exit(1)
    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        token = f.read().strip()
    if not token:
        logger.error("Токен пустой")
        sys.exit(1)
    return token


def read_admin_ids() -> List[int]:
    if not ADMIN_FILE.exists():
        logger.warning(f"Файл admin_ids.txt не найден: {ADMIN_FILE}")
        return []
    ids: List[int] = []
    try:
        with open(ADMIN_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and line.lstrip("-").isdigit():
                    ids.append(int(line))
    except Exception as e:
        logger.error(f"Ошибка чтения admin_ids.txt: {e}")
    return ids


TOKEN = read_token()
ADMIN_IDS = read_admin_ids()

# ------------------------- Состояния -------------------------
(
    MAIN_MENU,
    FILES_LIST,
    FILE_MENU,
    NEW_FILE_NAME,
    RENAME_FILE_NAME,
    CONFIRM_DELETE,
    BULK_PERCENT_INPUT,
    SINGLE_SELECT_ROW,
    SINGLE_SELECT_COL,
    SINGLE_SELECT_OP,
    SINGLE_VALUE_INPUT,
    UPLOAD_ACTION,
    UPLOAD_REPLACE_SELECT,
    UPLOAD_NEW_NAME,
) = range(14)


# ------------------------- Декоратор -------------------------
def admin_only(func):
    async def wrapper(update: Update, context: ContextTypes.DEFAULT_TYPE):
        user = update.effective_user
        if user is None:
            return
        if ADMIN_IDS and user.id not in ADMIN_IDS:
            if update.callback_query:
                await update.callback_query.answer("⛔ У вас нет прав.", show_alert=True)
                try:
                    await update.callback_query.edit_message_text("⛔ Доступ запрещён.")
                except Exception:
                    pass
            elif update.message:
                await update.message.reply_text("⛔ У вас нет прав на использование этого бота.")
            return
        return await func(update, context)
    return wrapper


# ------------------------- Вспомогательные функции -------------------------
CANCEL_KEYBOARD = InlineKeyboardMarkup(
    [[InlineKeyboardButton("❌ Отменить", callback_data="cancel_op")]]
)


def is_valid_name(name: str) -> bool:
    if not name or len(name) > 40:
        return False
    if name == ORIGINAL_NAME:
        return False
    return bool(re.fullmatch(r"[0-9a-zA-Zа-яА-ЯёЁ_\-]+", name))


def get_file_path(name: str) -> Path:
    return PARTNERS_DIR / f"{name}.xlsx"


def list_files() -> List[Path]:
    files = [f for f in PARTNERS_DIR.glob("*.xlsx") if f.is_file()]
    files.sort(key=lambda p: (p.stem != ORIGINAL_NAME, p.stem.lower()))
    return files


def safe_int(v) -> Optional[int]:
    if v is None or v == "":
        return None
    try:
        return int(round(float(v)))
    except (ValueError, TypeError):
        return None


def get_grid_info(path: Path):
    wb = load_workbook(path)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    row_labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    return wb, ws, headers, row_labels, ws.max_row, ws.max_column


async def edit_or_send(update: Update, context: ContextTypes.DEFAULT_TYPE,
                       text: str, keyboard: Optional[list], parse_mode: str = "Markdown"):
    markup = InlineKeyboardMarkup(keyboard) if keyboard else None
    if update.callback_query:
        try:
            await update.callback_query.edit_message_text(
                text, reply_markup=markup, parse_mode=parse_mode
            )
            return
        except Exception:
            pass
    if update.message:
        await update.message.reply_text(text, reply_markup=markup, parse_mode=parse_mode)
    else:
        await context.bot.send_message(
            chat_id=update.effective_chat.id,
            text=text,
            reply_markup=markup,
            parse_mode=parse_mode,
        )


def validate_uploaded_xlsx(new_path: Path, orig_path: Path) -> Tuple[bool, str]:
    try:
        wb1 = load_workbook(orig_path)
        ws1 = wb1.active
        wb2 = load_workbook(new_path)
        ws2 = wb2.active
    except Exception as e:
        return False, f"не удалось прочитать файл ({e})"

    if ws1.max_row != ws2.max_row:
        return False, (f"число строк не совпадает "
                       f"(ожидается {ws1.max_row}, получено {ws2.max_row})")
    if ws1.max_column != ws2.max_column:
        return False, (f"число столбцов не совпадает "
                       f"(ожидается {ws1.max_column}, получено {ws2.max_column})")

    for c in range(2, ws1.max_column + 1):
        v1 = ws1.cell(1, c).value
        v2 = ws2.cell(1, c).value
        if str(v1) != str(v2):
            return False, f"заголовок столбца {c} не совпадает ('{v1}' vs '{v2}')"

    for r in range(2, ws1.max_row + 1):
        v1 = ws1.cell(r, 1).value
        v2 = ws2.cell(r, 1).value
        if str(v1) != str(v2):
            return False, f"метка строки {r} не совпадает ('{v1}' vs '{v2}')"

    for r in range(2, ws1.max_row + 1):
        for c in range(2, ws1.max_column + 1):
            v = ws2.cell(r, c).value
            if v is None or v == "":
                return False, f"ячейка ({r},{c}) пуста"
            try:
                float(v)
            except (ValueError, TypeError):
                return False, f"ячейка ({r},{c}) не является числом"

    return True, "OK"


# ------------------------- Меню -------------------------
async def show_main_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = "🔧 *Панель управления файлами цен*\n\nВыберите действие:"
    keyboard = [
        [InlineKeyboardButton("📁 Файлы цен", callback_data="files_list")],
        [InlineKeyboardButton("➕ Создать копию из 000000", callback_data="create_copy")],
        [InlineKeyboardButton("📤 Как загрузить свой xlsx", callback_data="upload_hint")],
    ]
    await edit_or_send(update, context, text, keyboard)


async def show_files_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    files = list_files()
    text = "*Файлы цен:*\n\n"
    if not files:
        text += "📭 Пока пусто. Создайте копию 000000."
    else:
        for f in files:
            marker = "🔒" if f.stem == ORIGINAL_NAME else "📄"
            size_kb = f.stat().st_size / 1024
            text += f"{marker} `{f.stem}.xlsx` ({size_kb:.1f} KB)\n"

    keyboard = []
    for f in files:
        prefix = "🔒 " if f.stem == ORIGINAL_NAME else "📄 "
        keyboard.append([InlineKeyboardButton(
            f"{prefix}{f.stem}", callback_data=f"fo:{f.stem}"
        )])
    keyboard.append([InlineKeyboardButton("◀ Главное меню", callback_data="main_menu")])
    await edit_or_send(update, context, text, keyboard)


async def show_file_menu(update: Update, context: ContextTypes.DEFAULT_TYPE, name: str):
    context.user_data["current_file"] = name
    path = get_file_path(name)
    if not path.exists():
        await edit_or_send(update, context, "❌ Файл не найден.", None)
        return

    is_original = (name == ORIGINAL_NAME)
    size_kb = path.stat().st_size / 1024
    text = f"📄 *{name}.xlsx*\n\n"
    text += f"Размер: {size_kb:.1f} KB\n"
    if is_original:
        text += "🔒 Это исходный шаблон. Доступен только просмотр и скачивание."

    keyboard = [
        [InlineKeyboardButton("📥 Скачать", callback_data="file_download")],
    ]
    if not is_original:
        keyboard += [
            [InlineKeyboardButton("💰 Изменить все цены (%)", callback_data="file_bulk")],
            [InlineKeyboardButton("🎯 Изменить одну цену", callback_data="file_single")],
            [InlineKeyboardButton("📝 Переименовать", callback_data="file_rename")],
            [InlineKeyboardButton("🗑️ Удалить", callback_data="file_delete")],
        ]
    keyboard.append([InlineKeyboardButton("◀ К списку файлов", callback_data="files_list")])
    await edit_or_send(update, context, text, keyboard)


# ------------------------- Точки входа -------------------------
@admin_only
async def start_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await show_main_menu(update, context)
    return MAIN_MENU


async def cancel_op(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    await show_main_menu(update, context)
    return MAIN_MENU


async def main_menu_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.callback_query:
        await update.callback_query.answer()
    await show_main_menu(update, context)
    return MAIN_MENU


async def files_list_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await show_files_list(update, context)
    return FILES_LIST


async def file_open_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = query.data.split(":", 1)[1]
    await show_file_menu(update, context, name)
    return FILE_MENU


async def upload_hint_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    text = (
        "📤 *Загрузка своего xlsx-файла*\n\n"
        "Просто отправьте боту файл `.xlsx`. Бот проверит:\n"
        "• структура совпадает с `000000.xlsx` (заголовки и первый столбец),\n"
        "• все ценовые ячейки заполнены числами.\n\n"
        "Если проверка пройдёт, вы сможете заменить им любой созданный файл "
        "или создать новый."
    )
    keyboard = [[InlineKeyboardButton("◀ Главное меню", callback_data="main_menu")]]
    await edit_or_send(update, context, text, keyboard)
    return MAIN_MENU


# ------------------------- Создание копии -------------------------
async def create_copy_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    if not ORIGINAL_FILE.exists():
        await query.edit_message_text(
            f"❌ Шаблон не найден: {ORIGINAL_FILE}",
            reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("◀ Назад", callback_data="main_menu")]]),
        )
        return MAIN_MENU
    await query.edit_message_text(
        "Введите имя нового файла (без .xlsx).\n"
        "Разрешены: буквы (рус/англ), цифры, `_` и `-`.",
        reply_markup=CANCEL_KEYBOARD,
        parse_mode="Markdown",
    )
    return NEW_FILE_NAME


async def create_copy_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = update.message.text.strip()
    if not is_valid_name(name):
        await update.message.reply_text(
            "❌ Недопустимое имя. Разрешены буквы (рус/англ), цифры, `_` и `-`.\n"
            "Попробуйте ещё:",
            reply_markup=CANCEL_KEYBOARD,
            parse_mode="Markdown",
        )
        return NEW_FILE_NAME
    dst = get_file_path(name)
    if dst.exists():
        await update.message.reply_text(
            "❌ Файл с таким именем уже существует. Введите другое:",
            reply_markup=CANCEL_KEYBOARD,
        )
        return NEW_FILE_NAME
    try:
        shutil.copy(str(ORIGINAL_FILE), str(dst))
    except Exception as e:
        await update.message.reply_text(f"❌ Ошибка копирования: {e}")
        return ConversationHandler.END
    await update.message.reply_text(f"✅ Создан файл `{name}.xlsx`.", parse_mode="Markdown")
    context.user_data["current_file"] = name
    await show_file_menu(update, context, name)
    return FILE_MENU


# ------------------------- Файловое меню -------------------------
async def file_download(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if not name:
        await query.edit_message_text("❌ Файл не выбран.")
        return MAIN_MENU
    path = get_file_path(name)
    if not path.exists():
        await query.edit_message_text("❌ Файл не найден.")
        return FILE_MENU
    with open(path, "rb") as f:
        await context.bot.send_document(
            chat_id=update.effective_chat.id,
            document=f,
            filename=f"{name}.xlsx",
            caption=f"📄 {name}.xlsx",
        )
    return FILE_MENU


# --------- Переименование ---------
async def file_rename_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if not name or name == ORIGINAL_NAME:
        await query.edit_message_text("❌ Нельзя переименовать этот файл.")
        return FILE_MENU
    await query.edit_message_text(
        f"Текущее имя: `{name}.xlsx`\nВведите новое имя (без .xlsx):",
        reply_markup=CANCEL_KEYBOARD,
        parse_mode="Markdown",
    )
    return RENAME_FILE_NAME


async def file_rename_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    new_name = update.message.text.strip()
    old_name = context.user_data.get("current_file")
    if not old_name:
        await update.message.reply_text("❌ Файл не выбран.")
        return ConversationHandler.END
    if not is_valid_name(new_name):
        await update.message.reply_text(
            "❌ Недопустимое имя. Разрешены буквы (рус/англ), цифры, `_` и `-`:",
            reply_markup=CANCEL_KEYBOARD,
            parse_mode="Markdown",
        )
        return RENAME_FILE_NAME
    src = get_file_path(old_name)
    dst = get_file_path(new_name)
    if dst.exists():
        await update.message.reply_text(
            "❌ Файл с таким именем уже существует. Введите другое:",
            reply_markup=CANCEL_KEYBOARD,
        )
        return RENAME_FILE_NAME
    try:
        src.rename(dst)
    except Exception as e:
        await update.message.reply_text(f"❌ Ошибка переименования: {e}")
        return ConversationHandler.END
    context.user_data["current_file"] = new_name
    await update.message.reply_text(f"✅ Файл переименован в `{new_name}.xlsx`.", parse_mode="Markdown")
    await show_file_menu(update, context, new_name)
    return FILE_MENU


# --------- Удаление ---------
async def file_delete_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if not name or name == ORIGINAL_NAME:
        await query.edit_message_text("❌ Нельзя удалить этот файл.")
        return FILE_MENU
    keyboard = [
        [InlineKeyboardButton("✅ Да, удалить", callback_data="delete_yes")],
        [InlineKeyboardButton("❌ Нет", callback_data="delete_no")],
    ]
    await query.edit_message_text(
        f"Удалить файл `{name}.xlsx`?",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="Markdown",
    )
    return CONFIRM_DELETE


async def delete_yes(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if not name or name == ORIGINAL_NAME:
        await query.edit_message_text("❌ Нельзя удалить этот файл.")
        return FILES_LIST
    path = get_file_path(name)
    try:
        if path.exists():
            path.unlink()
        context.user_data.pop("current_file", None)
        await query.edit_message_text(f"✅ Файл `{name}.xlsx` удалён.", parse_mode="Markdown")
    except Exception as e:
        await query.edit_message_text(f"❌ Ошибка удаления: {e}")
    await show_files_list(update, context)
    return FILES_LIST


async def delete_no(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if name:
        await show_file_menu(update, context, name)
    else:
        await show_files_list(update, context)
    return FILE_MENU


# ------------------------- Массовое изменение цен -------------------------
async def file_bulk_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if not name or name == ORIGINAL_NAME:
        await query.edit_message_text("❌ Нельзя изменять этот файл.")
        return FILE_MENU
    await query.edit_message_text(
        "Введите процент изменения цен.\n"
        "Примеры: `10` (+10%), `-7.5` (−7.5%).\n"
        "Округление до целых.",
        reply_markup=CANCEL_KEYBOARD,
        parse_mode="Markdown",
    )
    return BULK_PERCENT_INPUT


async def bulk_apply(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip().replace(",", ".")
    try:
        pct = float(text)
    except ValueError:
        await update.message.reply_text(
            "❌ Введите число (например, 10 или -5):",
            reply_markup=CANCEL_KEYBOARD,
        )
        return BULK_PERCENT_INPUT

    name = context.user_data.get("current_file")
    path = get_file_path(name) if name else None
    if not path or not path.exists():
        await update.message.reply_text("❌ Файл не найден.")
        return ConversationHandler.END

    try:
        wb = load_workbook(path)
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            for c in range(2, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is None or v == "":
                    continue
                try:
                    old = float(v)
                except (ValueError, TypeError):
                    continue
                ws.cell(r, c).value = int(round(old * (1 + pct / 100.0)))
        wb.save(path)
    except Exception as e:
        await update.message.reply_text(f"❌ Ошибка: {e}")
        return ConversationHandler.END

    sign = "+" if pct >= 0 else ""
    await update.message.reply_text(f"✅ Все цены изменены на {sign}{pct}%.")
    await show_file_menu(update, context, name)
    return FILE_MENU


# ------------------------- Изменение одной цены -------------------------
async def file_single_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    name = context.user_data.get("current_file")
    if not name or name == ORIGINAL_NAME:
        await query.edit_message_text("❌ Нельзя изменять этот файл.")
        return FILE_MENU
    path = get_file_path(name)
    if not path.exists():
        await query.edit_message_text("❌ Файл не найден.")
        return FILE_MENU

    wb, ws, headers, row_labels, mr, mc = get_grid_info(path)
    keyboard = []
    for r in range(2, mr + 1):
        label = row_labels[r - 1]
        keyboard.append([InlineKeyboardButton(str(label), callback_data=f"srow:{r}")])
    keyboard.append([InlineKeyboardButton("❌ Отменить", callback_data="cancel_op")])
    await query.edit_message_text(
        f"*{name}.xlsx*\nВыберите размер (строку):",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="Markdown",
    )
    return SINGLE_SELECT_ROW


async def single_row_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    r = int(query.data.split(":", 1)[1])
    context.user_data["single_row"] = r

    name = context.user_data.get("current_file")
    path = get_file_path(name)
    wb, ws, headers, row_labels, mr, mc = get_grid_info(path)

    # Разложим столбцы по 2 в ряд
    col_buttons = []
    for c in range(2, mc + 1):
        col_buttons.append(InlineKeyboardButton(str(headers[c - 1]), callback_data=f"scol:{c}"))
    keyboard = [col_buttons[i:i + 2] for i in range(0, len(col_buttons), 2)]
    keyboard.append([InlineKeyboardButton("❌ Отменить", callback_data="cancel_op")])

    await query.edit_message_text(
        f"Строка: *{row_labels[r - 1]}*\nВыберите столбец (код):",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="Markdown",
    )
    return SINGLE_SELECT_COL


async def single_col_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    c = int(query.data.split(":", 1)[1])
    context.user_data["single_col"] = c

    name = context.user_data.get("current_file")
    path = get_file_path(name)
    wb, ws, headers, row_labels, mr, mc = get_grid_info(path)
    r = context.user_data["single_row"]
    old_value = ws.cell(r, c).value

    keyboard = [
        [InlineKeyboardButton("🔢 Установить", callback_data="sop:set")],
        [InlineKeyboardButton("✖️ Умножить", callback_data="sop:mul")],
        [InlineKeyboardButton("➕ Прибавить", callback_data="sop:add")],
        [InlineKeyboardButton("➖ Вычесть", callback_data="sop:sub")],
        [InlineKeyboardButton("❌ Отменить", callback_data="cancel_op")],
    ]
    await query.edit_message_text(
        f"Строка: *{row_labels[r - 1]}*\n"
        f"Столбец: *{headers[c - 1]}*\n"
        f"Текущее значение: *{old_value}*\n\n"
        f"Выберите операцию:",
        reply_markup=InlineKeyboardMarkup(keyboard),
        parse_mode="Markdown",
    )
    return SINGLE_SELECT_OP


async def single_op_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    op = query.data.split(":", 1)[1]
    context.user_data["single_op"] = op

    examples = {
        "set": "новое значение (целое или дробное, округлится)",
        "mul": "коэффициент (например, 1.1)",
        "add": "число, которое прибавить (может быть отрицательным)",
        "sub": "число, которое вычесть",
    }
    await query.edit_message_text(
        f"Введите {examples.get(op, 'значение')}:",
        reply_markup=CANCEL_KEYBOARD,
    )
    return SINGLE_VALUE_INPUT


async def single_value_apply(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip().replace(",", ".")
    try:
        value = float(text)
    except ValueError:
        await update.message.reply_text(
            "❌ Введите число:", reply_markup=CANCEL_KEYBOARD
        )
        return SINGLE_VALUE_INPUT

    name = context.user_data.get("current_file")
    r = context.user_data.get("single_row")
    c = context.user_data.get("single_col")
    op = context.user_data.get("single_op")
    path = get_file_path(name) if name else None
    if not path or not path.exists() or r is None or c is None or not op:
        await update.message.reply_text("❌ Потерян контекст. Начните заново.")
        return ConversationHandler.END

    try:
        wb = load_workbook(path)
        ws = wb.active
        old_raw = ws.cell(r, c).value
        try:
            old = float(old_raw)
        except (ValueError, TypeError):
            old = 0.0

        if op == "set":
            new = value
        elif op == "mul":
            new = old * value
        elif op == "add":
            new = old + value
        elif op == "sub":
            new = old - value
        else:
            await update.message.reply_text("❌ Неизвестная операция.")
            return ConversationHandler.END

        new_int = int(round(new))
        ws.cell(r, c).value = new_int
        wb.save(path)
    except Exception as e:
        await update.message.reply_text(f"❌ Ошибка: {e}")
        return ConversationHandler.END

    await update.message.reply_text(
        f"✅ Значение обновлено: было {int(round(old)) if old == int(old) else old}, стало {new_int}."
    )
    await show_file_menu(update, context, name)
    return FILE_MENU


# ------------------------- Загрузка xlsx -------------------------
async def handle_uploaded_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if not doc or not doc.file_name or not doc.file_name.lower().endswith(".xlsx"):
        await update.message.reply_text("❌ Поддерживаются только .xlsx файлы.")
        return ConversationHandler.END

    if not ORIGINAL_FILE.exists():
        await update.message.reply_text(
            f"❌ Шаблон 000000.xlsx не найден ({ORIGINAL_FILE}). "
            f"Невозможно проверить структуру."
        )
        return ConversationHandler.END

    tmp_dir = Path(tempfile.mkdtemp(prefix="upload_"))
    tmp_path = tmp_dir / doc.file_name

    try:
        file = await doc.get_file()
        await file.download_to_drive(str(tmp_path))
    except Exception as e:
        await update.message.reply_text(f"❌ Ошибка скачивания файла: {e}")
        return ConversationHandler.END

    ok, msg = validate_uploaded_xlsx(tmp_path, ORIGINAL_FILE)
    if not ok:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        await update.message.reply_text(f"❌ Файл не прошёл проверку: {msg}")
        return ConversationHandler.END

    context.user_data["uploaded_path"] = str(tmp_path)
    context.user_data["uploaded_name"] = doc.file_name

    keyboard = [
        [InlineKeyboardButton("📝 Создать новый файл", callback_data="up_new")],
        [InlineKeyboardButton("♻️ Заменить существующий", callback_data="up_replace")],
        [InlineKeyboardButton("❌ Отменить", callback_data="cancel_op")],
    ]
    await update.message.reply_text(
        "✅ Файл проверен и совпадает по структуре с 000000.xlsx.\nЧто с ним сделать?",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    return UPLOAD_ACTION


async def upload_choose_new(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text(
        "Введите имя нового файла (без .xlsx):",
        reply_markup=CANCEL_KEYBOARD,
    )
    return UPLOAD_NEW_NAME


async def upload_new_name_input(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = update.message.text.strip()
    if not is_valid_name(name):
        await update.message.reply_text(
            "❌ Недопустимое имя. Разрешены буквы (рус/англ), цифры, `_` и `-`:",
            reply_markup=CANCEL_KEYBOARD,
            parse_mode="Markdown",
        )
        return UPLOAD_NEW_NAME
    dst = get_file_path(name)
    if dst.exists():
        await update.message.reply_text(
            "❌ Файл с таким именем уже существует. Введите другое:",
            reply_markup=CANCEL_KEYBOARD,
        )
        return UPLOAD_NEW_NAME

    src_str = context.user_data.pop("uploaded_path", None)
    if not src_str:
        await update.message.reply_text("❌ Временный файл не найден.")
        return ConversationHandler.END
    src = Path(src_str)
    if not src.exists():
        await update.message.reply_text("❌ Временный файл не найден.")
        return ConversationHandler.END

    try:
        shutil.copy(str(src), str(dst))
    except Exception as e:
        await update.message.reply_text(f"❌ Ошибка сохранения: {e}")
        return ConversationHandler.END
    finally:
        shutil.rmtree(src.parent, ignore_errors=True)

    await update.message.reply_text(f"✅ Создан файл `{name}.xlsx`.", parse_mode="Markdown")
    context.user_data["current_file"] = name
    await show_file_menu(update, context, name)
    return FILE_MENU


async def upload_choose_replace(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    files = [f for f in list_files() if f.stem != ORIGINAL_NAME]
    if not files:
        await query.edit_message_text(
            "❌ Нет созданных файлов для замены (кроме 000000).",
            reply_markup=InlineKeyboardMarkup(
                [[InlineKeyboardButton("◀ Главное меню", callback_data="main_menu")]]
            ),
        )
        return MAIN_MENU

    context.user_data["upload_replace_files"] = [f.stem for f in files]
    keyboard = []
    for i, f in enumerate(files):
        keyboard.append([InlineKeyboardButton(
            f"📄 {f.stem}", callback_data=f"up_repl_idx:{i}"
        )])
    keyboard.append([InlineKeyboardButton("❌ Отменить", callback_data="cancel_op")])
    await query.edit_message_text(
        "Выберите файл для замены:",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    return UPLOAD_REPLACE_SELECT


async def upload_replace_execute(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    idx = int(query.data.split(":", 1)[1])
    names = context.user_data.get("upload_replace_files") or []
    if idx < 0 or idx >= len(names):
        await query.edit_message_text("❌ Некорректный выбор.")
        return MAIN_MENU
    target_name = names[idx]

    src_str = context.user_data.pop("uploaded_path", None)
    if not src_str:
        await query.edit_message_text("❌ Временный файл не найден.")
        return MAIN_MENU
    src = Path(src_str)
    if not src.exists():
        await query.edit_message_text("❌ Временный файл не найден.")
        return MAIN_MENU

    dst = get_file_path(target_name)
    try:
        shutil.copy(str(src), str(dst))
    except Exception as e:
        await query.edit_message_text(f"❌ Ошибка замены: {e}")
        shutil.rmtree(src.parent, ignore_errors=True)
        return MAIN_MENU
    finally:
        shutil.rmtree(src.parent, ignore_errors=True)

    context.user_data["current_file"] = target_name
    await query.edit_message_text(
        f"✅ Файл `{target_name}.xlsx` заменён.", parse_mode="Markdown"
    )
    await show_file_menu(update, context, target_name)
    return FILE_MENU


# ------------------------- Регистрация -------------------------
def register_handlers(app: Application):
    conv = ConversationHandler(
        entry_points=[
            CommandHandler("start", start_cmd),
            MessageHandler(filters.Document.FileExtension("xlsx"), handle_uploaded_document),
        ],
        states={
            MAIN_MENU: [
                CallbackQueryHandler(files_list_cb, pattern="^files_list$"),
                CallbackQueryHandler(create_copy_start, pattern="^create_copy$"),
                CallbackQueryHandler(upload_hint_cb, pattern="^upload_hint$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
                CallbackQueryHandler(main_menu_cb, pattern="^main_menu$"),
            ],
            FILES_LIST: [
                CallbackQueryHandler(file_open_cb, pattern=r"^fo:.+$"),
                CallbackQueryHandler(main_menu_cb, pattern="^main_menu$"),
                CallbackQueryHandler(files_list_cb, pattern="^files_list$"),
            ],
            FILE_MENU: [
                CallbackQueryHandler(file_download, pattern="^file_download$"),
                CallbackQueryHandler(file_bulk_start, pattern="^file_bulk$"),
                CallbackQueryHandler(file_single_start, pattern="^file_single$"),
                CallbackQueryHandler(file_rename_start, pattern="^file_rename$"),
                CallbackQueryHandler(file_delete_start, pattern="^file_delete$"),
                CallbackQueryHandler(files_list_cb, pattern="^files_list$"),
                CallbackQueryHandler(main_menu_cb, pattern="^main_menu$"),
            ],
            NEW_FILE_NAME: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, create_copy_input),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            RENAME_FILE_NAME: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, file_rename_input),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            CONFIRM_DELETE: [
                CallbackQueryHandler(delete_yes, pattern="^delete_yes$"),
                CallbackQueryHandler(delete_no, pattern="^delete_no$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            BULK_PERCENT_INPUT: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, bulk_apply),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            SINGLE_SELECT_ROW: [
                CallbackQueryHandler(single_row_selected, pattern=r"^srow:\d+$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            SINGLE_SELECT_COL: [
                CallbackQueryHandler(single_col_selected, pattern=r"^scol:\d+$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            SINGLE_SELECT_OP: [
                CallbackQueryHandler(single_op_selected, pattern=r"^sop:(set|mul|add|sub)$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            SINGLE_VALUE_INPUT: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, single_value_apply),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            UPLOAD_ACTION: [
                CallbackQueryHandler(upload_choose_new, pattern="^up_new$"),
                CallbackQueryHandler(upload_choose_replace, pattern="^up_replace$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            UPLOAD_REPLACE_SELECT: [
                CallbackQueryHandler(upload_replace_execute, pattern=r"^up_repl_idx:\d+$"),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
            UPLOAD_NEW_NAME: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, upload_new_name_input),
                CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            ],
        },
        fallbacks=[
            CommandHandler("cancel", cancel_op),
            CallbackQueryHandler(cancel_op, pattern="^cancel_op$"),
            CallbackQueryHandler(main_menu_cb, pattern="^main_menu$"),
            # Разрешаем прислать файл в любой момент
            MessageHandler(filters.Document.FileExtension("xlsx"), handle_uploaded_document),
        ],
        per_message=False,
        allow_reentry=True,
    )
    app.add_handler(conv)


# ------------------------- Запуск -------------------------
async def main():
    if not ORIGINAL_FILE.exists():
        logger.warning(f"ВНИМАНИЕ: шаблон {ORIGINAL_FILE} не найден. "
                       f"Некоторые функции будут недоступны.")

    request = HTTPXRequest(
        connect_timeout=30.0,
        read_timeout=120.0,
        write_timeout=30.0,
        pool_timeout=30.0,
    )
    application = Application.builder().token(TOKEN).request(request).build()
    register_handlers(application)

    await application.initialize()
    await application.start()
    logger.info("Бот для управления ценами партнёров запущен.")
    await application.updater.start_polling()
    await asyncio.Event().wait()


if __name__ == "__main__":
    import asyncio
    asyncio.run(main())